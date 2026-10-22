#!/usr/bin/env python3
"""Convertit la note éditoriale Markdown en .docx propre, paginé."""
import re
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = sys.argv[1]
DST = sys.argv[2]

BODY_FONT = "Cambria"
ACCENT = RGBColor(0x1a, 0x1a, 0x1a)

doc = Document()
cp = doc.core_properties
cp.title = "Note sur la Bible de projet v6.0"
cp.subject = "CE QUE JE TE DOIS"
cp.author = "Note éditoriale"
cp.comments = "Document de travail établi à la demande de l'autrice."

# --- page setup
for section in doc.sections:
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.6)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = BODY_FONT
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12
normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

for name, size, bold, before, after in (
    ("Heading 1", 15, True, 18, 8),
    ("Heading 2", 12, True, 14, 6),
    ("Heading 3", 11, True, 10, 4),
):
    st = styles[name]
    st.font.name = BODY_FONT
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.color.rgb = ACCENT
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True


def add_page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for instr in ('begin', 'PAGE', 'end'):
        fld = OxmlElement('w:fldChar')
        fld.set(qn('w:fldCharType'), instr)
        run._r.append(fld)
    p.runs[0].font.size = Pt(8.5)
    p.runs[0].font.name = BODY_FONT


add_page_number_footer(doc.sections[0])

# --- inline formatting
INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|_[^_]+?_)")


def add_runs(par, text, base_bold=False, base_italic=False, size=None):
    text = text.replace("’", "’")
    for part in INLINE.split(text):
        if not part:
            continue
        bold, italic, mono = base_bold, base_italic, False
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, bold = part[2:-2], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, italic = part[1:-1], True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            part, mono = part[1:-1], True
        run = par.add_run(part)
        run.bold = bold
        run.italic = italic
        if mono:
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif size:
            run.font.size = Pt(size)


def parse_row(line):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


def add_table(rows, header=True):
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        if header and i == 0:
            tr = table.rows[0]._tr
            trPr = tr.get_or_add_trPr()
            tblHeader = OxmlElement('w:tblHeader')
            tblHeader.set(qn('w:val'), 'true')
            trPr.append(tblHeader)
            for cell in table.rows[0].cells:
                tcPr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'), 'EDEDED')
                tcPr.append(shd)
        for j in range(ncols):
            txt = row[j] if j < len(row) else ""
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            add_runs(p, txt, base_bold=(header and i == 0), size=9.5)
    doc.add_paragraph()


# --- front page
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(120)
r = p.add_run("CE QUE JE TE DOIS")
r.font.size = Pt(12)
r.font.name = BODY_FONT
r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(24)
r = p.add_run("Note sur la Bible de projet v6.0")
r.font.size = Pt(26)
r.font.name = BODY_FONT
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

p = doc.add_paragraph()
r = p.add_run("Ce qui bloque, et ce qu'une v7 doit contenir")
r.font.size = Pt(14)
r.italic = True
r.font.name = BODY_FONT
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
pPr = p._p.get_or_add_pPr()
bdr = OxmlElement('w:pBdr')
bottom = OxmlElement('w:bottom')
bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
bottom.set(qn('w:space'), '8'); bottom.set(qn('w:color'), '999999')
bdr.append(bottom); pPr.append(bdr)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(56)
r = p.add_run("Document de travail établi à la demande de l'autrice\npour transmission à son éditeur")
r.font.size = Pt(10.5)
r.font.name = BODY_FONT
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(28)
r = p.add_run("22 octobre 2026")
r.font.size = Pt(10.5)
r.font.name = BODY_FONT
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# --- body
_front = True
lines = open(SRC, encoding="utf-8").read().split("\n")
i = 0
first_h1_skipped = False
while i < len(lines):
    line = lines[i]
    s = line.strip()

    if not s:
        i += 1
        continue

    if s.startswith("# "):
        i += 1
        continue  # titre de la page de garde

    if s.startswith("### "):
        doc.add_heading(s[4:], level=2)
        i += 1
        continue
    if s.startswith("## "):
        if _front:
            _front = False
            i += 1
            continue
        doc.add_heading(s[3:], level=1)
        i += 1
        continue

    if set(s) <= {"-", " "} and len(s) >= 3:
        i += 1
        continue

    if s.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            raw = lines[i].strip()
            if not re.match(r"^\|[\s:\-|]+\|$", raw):
                rows.append(parse_row(raw))
            i += 1
        if rows:
            add_table(rows)
        continue

    if s.startswith("> "):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.right_indent = Cm(0.6)
        add_runs(p, s[2:], base_italic=True)
        i += 1
        continue

    m = re.match(r"^(\d+)\.\s+(.*)$", s)
    if m:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.85)
        p.paragraph_format.first_line_indent = Cm(-0.85)
        p.paragraph_format.space_after = Pt(4)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        num = p.add_run(m.group(1) + ".\t")
        num.bold = False
        add_runs(p, m.group(2))
        i += 1
        continue

    if s.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        add_runs(p, s[2:])
        i += 1
        continue

    # paragraphe (fusionner les lignes de continuation sans puce)
    buf = [s]
    i += 1
    while i < len(lines):
        nxt = lines[i].strip()
        if (not nxt or nxt.startswith(("#", "|", "- ", "> "))
                or re.match(r"^\d+\.\s", nxt) or set(nxt) <= {"-", " "}):
            break
        buf.append(nxt)
        i += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_runs(p, " ".join(buf))

doc.save(DST)
print(f"écrit : {DST}")
